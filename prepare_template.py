"""Create a reusable blank template without retaining historical statistics."""
import argparse
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.packaging.core import DocumentProperties


def prepare(source, target):
    source, target = Path(source), Path(target)
    if source.resolve() == target.resolve() or target.exists():
        raise ValueError('Choose a new destination; the source must be preserved')
    workbook = openpyxl.load_workbook(source)
    for sheet in list(workbook):
        if 'AI 제거버전' not in sheet.title:
            workbook.remove(sheet)
            continue
        for start in (5, 43, 81, 119, 157):
            for row in sheet.iter_rows(min_row=start, max_row=start + 31, min_col=2, max_col=27):
                for cell in row:
                    cell.value = None
        for row in sheet:
            for cell in row:
                if not isinstance(cell, MergedCell):
                    cell.comment = None
                    cell.hyperlink = None
    if not workbook.sheetnames:
        raise ValueError('No supported AI-removal sheet')
    workbook.properties = DocumentProperties(creator='PACS statistics', title='Blank statistics template')
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source')
    parser.add_argument('target')
    args = parser.parse_args()
    prepare(args.source, args.target)
