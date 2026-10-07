"""Template-preserving statistics core; PACS screen capture is not implemented yet."""
from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

KOREA = ZoneInfo('Asia/Seoul')
STARTS = (5, 43, 81, 119, 157)


def aggregate(records, day, ae_title):
    """One record per study, with a stable identifier supplied by the collector."""
    buckets = defaultdict(lambda: [0, 0, 0, 0, 0])
    seen = set()
    for item in records:
        timestamp = datetime.fromisoformat(item['study_datetime'])
        if timestamp.tzinfo:
            timestamp = timestamp.astimezone(KOREA)
        if timestamp.date() != day or item['source_aetitle'].strip() != ae_title:
            continue
        study_id = str(item['study_id']).strip()
        if not study_id:
            raise ValueError('Missing stable study identifier')
        if study_id in seen:
            raise ValueError('Duplicate study identifier; review screen collection')
        seen.add(study_id)
        series, images = item['series_count'], item['instance_count']
        if type(series) is not int or series < 1 or type(images) is not int or images < 1:
            raise ValueError('Invalid series or instance count')
        modalities = set(re.split(r'[\\/,;\s]+', item['modality'].upper().strip()))
        # Explicit classification is required; do not infer Chest from arbitrary text.
        if type(item['is_chest']) is not bool:
            raise ValueError('Chest classification must be a boolean')
        deduction = 0
        if item['is_chest'] and {'DX', 'PR'} <= modalities:
            if series not in (2, 3, 4):
                raise ValueError('Unsupported Chest AI series count; review required')
            deduction = series - 1
            if images - deduction < 1:
                raise ValueError('AI deduction leaves no radiograph; review required')
        values = buckets[timestamp.hour // 2]
        values[0] += 1
        values[1] += images
        if deduction:
            values[deduction + 1] += 1
    return buckets


def export_day(template, destination, records, day, ae_title, previous=None):
    template, destination = Path(template), Path(destination)
    previous = Path(previous) if previous else None
    if destination.resolve() in {template.resolve(), previous.resolve() if previous else template.resolve()}:
        raise ValueError('Save to a new file; never overwrite the template or previous output')
    if destination.exists():
        raise ValueError('Destination already exists')
    if not re.fullmatch(r'[A-Za-z0-9_-]+', ae_title):
        raise ValueError('AE Title must contain letters, numbers, underscores or hyphens')
    buckets = aggregate(records, day, ae_title)
    month = day.strftime('%Y%m')
    sheet_name = f'{month}_{ae_title}'
    workbook = openpyxl.load_workbook(previous or template)
    if previous:
        if workbook.sheetnames != [sheet_name]:
            raise ValueError('Previous output must match this month and AE Title')
        sheet = workbook[sheet_name]
        if sheet['T2'].value != f'AETitle : {ae_title}':
            raise ValueError('Previous output AE Title does not match')
    else:
        source = f'{month} AI 제거버전'
        if source not in workbook:
            raise ValueError('No AI-removal template for this month; template expansion is required')
        sheet = workbook[source]
        for other in list(workbook):
            if other is not sheet:
                workbook.remove(other)
        sheet.title = sheet_name
        # Historical data in the example template belongs to ER_DR, not a new report.
        for start in STARTS:
            for row in sheet.iter_rows(min_row=start, max_row=start + 30, min_col=2, max_col=25):
                for cell in row:
                    cell.value = None
        sheet['T2'] = f'AETitle : {ae_title}'
    for slot in range(12):
        exam_col, image_col = 2 + slot * 2, 3 + slot * 2
        exam, images, srs2, srs3, srs4 = buckets[slot]
        corrected = images - srs2 - srs3 * 2 - srs4 * 3
        for start, exams, count in [(5, exam, corrected), (43, exam, images),
                                     (81, None, srs2), (119, None, srs3), (157, None, srs4)]:
            row = start + day.day - 1
            sheet.cell(row, exam_col, exams).value = exams
            sheet.cell(row, image_col, count)
    # Materialize totals and final cells so Excel need not recalculate formula caches.
    # This also removes the sample workbook's incorrect cross-date references.
    for day_index in range(31):
        for slot in range(12):
            ec, ic = 2 + 2 * slot, 3 + 2 * slot
            raw = sheet.cell(43 + day_index, ic).value
            if raw is not None:
                counts = [sheet.cell(start + day_index, ic).value or 0 for start in (81, 119, 157)]
                corrected = raw - sum(n * weight for n, weight in zip(counts, (1, 2, 3)))
                if corrected < 0:
                    raise ValueError('Previous data has negative corrected image count')
                sheet.cell(5 + day_index, ic, corrected)
                sheet.cell(5 + day_index, ec).value = sheet.cell(43 + day_index, ec).value
    for start in STARTS:
        for row in range(start, start + 31):
            for col, columns in [(26, range(2, 26, 2)), (27, range(3, 26, 2))]:
                vals = [sheet.cell(row, c).value for c in columns]
                sheet.cell(row, col).value = sum(v for v in vals if v is not None) if any(v is not None for v in vals) else None
        for col in range(2, 28):
            vals = [sheet.cell(row, col).value for row in range(start, start + 31)]
            sheet.cell(start + 31, col).value = sum(v for v in vals if v is not None) if any(v is not None for v in vals) else None
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return sum(v[0] for v in buckets.values())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--template', required=True)
    parser.add_argument('--records', required=True, help='Normalized JSON; not a PACS export')
    parser.add_argument('--aetitle', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--previous', help='Prior generated output for the same month and AE Title')
    parser.add_argument('--date', type=date.fromisoformat, default=datetime.now(KOREA).date())
    args = parser.parse_args()
    records = json.loads(Path(args.records).read_text(encoding='utf-8'))
    count = export_day(args.template, args.output, records, args.date, args.aetitle, args.previous)
    print(f'Saved {args.output}: {count} studies for {args.date} / {args.aetitle}')


if __name__ == '__main__':
    main()
